"""Genera docs/Informe.pdf desde docs/informe.md (requiere pandoc, python-docx y libreoffice)."""
import os, subprocess
from docx import Document
from docx.shared import Pt, Cm

d = os.path.dirname(os.path.abspath(__file__))
md, docx, = os.path.join(d, "informe.md"), os.path.join(d, "Informe.docx")
subprocess.run(["pandoc", md, "-o", docx], check=True)
doc = Document(docx)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(1.8); s.top_margin = s.bottom_margin = Cm(1.6)
for st in doc.styles:
    try:
        if st.font is not None:
            st.font.name = "Calibri"
            st.font.size = Pt({"Heading 1": 13, "Title": 16}.get(st.name, 10.5))
    except Exception:
        pass
for p in doc.paragraphs:
    p.paragraph_format.space_after = Pt(3); p.paragraph_format.space_before = Pt(0)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ANCHOS = {4: [1.4, 1.2, 7.4, 7.4], 3: [3.4, 5.4, 8.6]}
for t in doc.tables:
    t.autofit = False
    tblPr = t._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "999999")
        b.append(e)
    tblPr.append(b)
    anchos = ANCHOS.get(len(t.columns))
    if anchos:
        for j, col in enumerate(t.columns):
            col.width = Cm(anchos[j])
    for i, row in enumerate(t.rows):
        for j, c in enumerate(row.cells):
            if anchos:
                c.width = Cm(anchos[j])
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    if i == 0: r.font.bold = True
doc.save(docx)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", d, docx], check=True, capture_output=True)
os.remove(docx)
print(subprocess.run(["pdfinfo", os.path.join(d, "Informe.pdf")], capture_output=True, text=True).stdout)
